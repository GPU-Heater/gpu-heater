import os
import PyPDF2
import xlrd
import openpyxl
import docx
from pptx import Presentation
from bs4 import BeautifulSoup
from blueprints.config import DEFAULT_EXTS, TEXT_EXTS

def parse_document(file_path, ext):
    file_text_content = ""
    try:
        if ext in ['.html', '.htm', '.xhtml']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                soup = BeautifulSoup(f.read(), 'html.parser')
                for tag in soup(["script", "style"]):
                    tag.decompose()
                file_text_content = soup.get_text(separator='\n', strip=True)

        elif ext == '.pdf':
            with open(file_path, 'rb') as f:
                pdf = PyPDF2.PdfReader(f)
                file_text_content = "\n".join([p.extract_text() for p in pdf.pages if p.extract_text()])

        elif ext == '.docx':
            doc = docx.Document(file_path)
            file_text_content = "\n".join([para.text for para in doc.paragraphs])

        elif ext == '.pptx':
            prs = Presentation(file_path)
            text_runs = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_runs.append(shape.text)
            file_text_content = "\n".join(text_runs)

        elif ext in ['.xls', '.xlsx']:
            excel_text = []
            if ext == '.xlsx':
                wb = openpyxl.load_workbook(file_path, data_only=True)
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    excel_text.append(f"\n--- Sheet: {sheet_name} ---")
                    for row in sheet.iter_rows(values_only=True):
                        row_vals = [str(cell) if cell is not None else "" for cell in row]
                        if any(row_vals):
                            excel_text.append(",".join(row_vals))
            elif ext == '.xls':
                wb = xlrd.open_workbook(file_path)
                for sheet_idx in range(wb.nsheets):
                    sheet = wb.sheet_by_index(sheet_idx)
                    excel_text.append(f"\n--- Sheet: {sheet.name} ---")
                    for row_idx in range(sheet.nrows):
                        row_vals = [str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
                        if any(row_vals):
                            excel_text.append(",".join(row_vals))
            file_text_content = "\n".join(excel_text)

        else:
            with open(file_path, 'rb') as f:
                chunk = f.read(1024)
                if b'\x00' in chunk:
                    return ""

            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                file_text_content = f.read()

    except Exception as e:
        file_text_content = f"[File Read Error: {e}]"
        
    return file_text_content